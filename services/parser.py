"""Text extraction from PDF and DOCX documents."""

from __future__ import annotations

from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from services.exceptions import (
    DocumentParsingError,
    EmptyDocumentError,
    UnsupportedFileTypeError,
)
from utils.logger import get_logger

logger = get_logger(__name__)

SUPPORTED_EXTENSIONS = (".pdf", ".docx")


class DocumentParser:
    """Extracts plain text from the document formats the app supports."""

    def parse(self, file_path: str | Path) -> str:
        """Extract the text content of a document.

        Args:
            file_path: Path to a ``.pdf`` or ``.docx`` file.

        Returns:
            The extracted text.

        Raises:
            FileNotFoundError: If the path does not exist.
            UnsupportedFileTypeError: If the extension is not supported.
            DocumentParsingError: If the file is corrupt or unreadable.
            EmptyDocumentError: If no text could be extracted.
        """
        path = Path(file_path)

        if not path.is_file():
            raise FileNotFoundError(f"Document not found: {path}")

        extension = path.suffix.lower()

        if extension == ".pdf":
            text = self._parse_pdf(path)
        elif extension == ".docx":
            text = self._parse_docx(path)
        else:
            raise UnsupportedFileTypeError(
                f"Unsupported file type {extension!r}. "
                f"Supported types: {', '.join(SUPPORTED_EXTENSIONS)}."
            )

        text = text.strip()

        if not text:
            raise EmptyDocumentError(
                f"No text could be extracted from {path.name}. "
                "The file may be a scanned image and would need OCR."
            )

        logger.info("Parsed %s (%d characters)", path.name, len(text))
        return text

    def _parse_pdf(self, path: Path) -> str:
        """Extract text from a PDF, page by page."""
        pages: list[str] = []
        try:
            with fitz.open(path) as pdf:
                for page in pdf:
                    pages.append(page.get_text("text"))
        except Exception as exc:  # PyMuPDF raises a variety of low-level errors
            raise DocumentParsingError(f"Could not read the PDF {path.name}: {exc}") from exc

        return "\n".join(pages)

    def _parse_docx(self, path: Path) -> str:
        """Extract paragraph and table text from a DOCX file.

        Legal documents frequently keep obligations inside tables, so table
        cells are extracted too rather than silently dropped.
        """
        try:
            document = Document(str(path))
        except PackageNotFoundError as exc:
            raise DocumentParsingError(
                f"{path.name} is not a valid DOCX file (legacy .doc is not supported)."
            ) from exc
        except Exception as exc:
            raise DocumentParsingError(f"Could not read the DOCX {path.name}: {exc}") from exc

        blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    blocks.append(" | ".join(cells))

        return "\n".join(blocks)
